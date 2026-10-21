"""
Build one district-month table from D2 (cb_t3T-district yearly workbooks) in the D1 (FI_CB_095_S2) column
layout, append D1 (Aug 2569, scraped from the BOT web table), and write tidy CSVs.

Usage: python build_district.py <d2_xlsx_dir> <d1_rows_txt> <d1_na_txt> <population_csv> <out_dir>
"""
import sys, os, re, glob
import openpyxl, pandas as pd
sys.path.insert(0, os.path.dirname(__file__))
from parse_bot import parse_period, num

D1_COLS = ['offices', 'dep_demand', 'dep_savings', 'dep_time', 'dep_pn', 'dep_ncd', 'dep_total',
           'loan_overdraft', 'loan_loans', 'loan_bills_other', 'loan_total']
D1_TH = ['จำนวนสำนักงาน', 'จ่ายคืนเมื่อทวงถาม', 'ออมทรัพย์', 'ประจำ', 'ตั๋วสัญญาใช้เงิน', 'บัตรเงินฝาก', 'รวมเงินฝาก',
         'เงินเบิกเกินบัญชี', 'เงินให้กู้', 'ตั๋วเงินและอื่นๆ', 'รวมเงินให้สินเชื่อ']


def clean(v):
    return re.sub(r'\s+', '', str(v)) if v is not None else ''


def header_map(rows):
    """Find the field-name header row and map column index -> D1 field (bills/other may map from 2 columns)."""
    hi = next(i for i, r in enumerate(rows[:15]) if any('ออมทรัพย์' in clean(v) for v in r))
    hdr = [clean(v) for v in rows[hi]]
    m, seen_total, in_loans = {1: 'offices'}, 0, False
    for j, h in enumerate(hdr):
        if j < 2 or not h:
            continue
        if h == 'รวม':
            m[j] = 'dep_total' if seen_total == 0 else 'loan_total'
            seen_total += 1; in_loans = seen_total >= 1; continue
        if 'กระแสรายวัน' in h or 'ทวงถาม' in h: m[j] = 'dep_demand'
        elif 'ออมทรัพย์' in h: m[j] = 'dep_savings'
        elif 'ประจำ' in h or 'สิ้นระยะเวลา' in h: m[j] = 'dep_time'
        elif 'ตั๋วสัญญาใช้เงิน' in h: m[j] = 'dep_pn'
        elif 'บัตรเงินฝาก' in h: m[j] = 'dep_ncd'
        elif 'เบิกเกินบัญชี' in h: m[j] = 'loan_overdraft'
        elif 'เงินให้กู้' in h: m[j] = 'loan_loans'
        elif 'ตั๋วเงิน' in h: m[j] = 'loan_bills_other'
        elif 'อื่น' in h: m[j] = 'loan_bills_other' if in_loans else 'dep_other'
        else: raise ValueError(f'unknown header {h!r}')
    return hi, m


def strip_prefix(name):
    return re.sub(r'^(อำเภอ|เขต|กิ่งอำเภอ)\s*', '', name.replace('*', '')).strip()


def read_d2(d2_dir):
    recs, log = [], []
    for f in sorted(glob.glob(os.path.join(d2_dir, '*.xlsx'))):
        fname = os.path.basename(f)
        loan_basis = 'S2' if '_S2' in fname else 'old'
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for sn in wb.sheetnames:
            p = parse_period(sn)
            if not p:
                continue
            rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
            hi, cmap = header_map(rows)
            log.append((fname, sn, p, sorted(set(cmap.values()))))
            prov = None
            for r in rows[hi + 1:]:
                k = str(r[0] or '').strip()
                if k.startswith('จังหวัด'):
                    prov = k.replace('จังหวัด', '').strip(); continue
                mm = re.match(r'^(\d{6})\s+(.*)$', k)
                is_hq = 'สำนักงานใหญ่' in k
                if not (mm or is_hq):
                    continue
                rec = {'year_be': p[0], 'month': p[1], 'source': 'D2', 'source_file': fname, 'source_sheet': sn,
                       'loan_basis': loan_basis, 'is_hq': is_hq, 'is_suppressed': False,
                       'amphoe_code': 'HQ' if is_hq else mm.group(1)[:4],
                       'amphoe_name': 'สำนักงานใหญ่' if is_hq else strip_prefix(mm.group(2)),
                       'prov_name': None if is_hq else prov}
                for j, col in cmap.items():
                    v = num(r[j]) if j < len(r) else None
                    if col == 'loan_bills_other' and col in rec and rec[col] is not None:
                        rec[col] = rec[col] + (v or 0)      # old layout: ตั๋วเงิน + อื่นๆ -> ตั๋วเงินและอื่นๆ
                    else:
                        rec[col] = v
                recs.append(rec)
    return pd.DataFrame(recs), log


def read_d1(rows_txt, na_txt):
    recs = []
    for line in open(rows_txt, encoding='utf-8').read().strip().split(';'):
        p = line.split(',')
        recs.append({'amphoe_code': p[0], **{c: num(v) for c, v in zip(D1_COLS, p[1:])}, 'is_suppressed': False})
    for c in open(na_txt, encoding='utf-8').read().strip().split(','):
        recs.append({'amphoe_code': c, 'is_suppressed': True})
    d = pd.DataFrame(recs)
    d['year_be'], d['month'] = 2569, 8
    d['source'], d['source_file'], d['source_sheet'], d['loan_basis'], d['is_hq'] = \
        'D1', 'FI_CB_095_S2 (web table)', 'ส.ค. 2569 p', 'S2', False
    return d


def main(d2_dir, d1_rows, d1_na, pop_csv, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    d2, log = read_d2(d2_dir)

    # duplicate months: 2567 & 2567_S2 (Jul-Dec 67), 2568 & 2568_S2 (Jan-Jun 68). Keep S2 in the main table.
    overlap = d2.groupby(['year_be', 'month']).loan_basis.nunique()
    overlap = overlap[overlap > 1].index
    d2['in_overlap'] = d2.set_index(['year_be', 'month']).index.isin(overlap)
    d2[d2.in_overlap].to_csv(f'{out_dir}/d2_overlap_old_vs_S2.csv', index=False, encoding='utf-8-sig')
    d2 = d2[~(d2.in_overlap & (d2.loan_basis == 'old'))].drop(columns='in_overlap')

    # add rows for districts hidden in D2 so the layout matches D1 (all districts listed, hidden = suppressed)
    pop = pd.read_csv(pop_csv, dtype={'amphoe_code': str, 'prov_code': str}, encoding='utf-8-sig')
    names = pop.sort_values('year_be').groupby('amphoe_code').amphoe_name.last().map(strip_prefix)
    universe = pop[['year_be', 'amphoe_code']]
    months = d2[['year_be', 'month']].drop_duplicates()
    months['pop_year'] = months.year_be.clip(upper=pop.year_be.max())
    full = months.merge(universe.rename(columns={'year_be': 'pop_year'}), on='pop_year').drop(columns='pop_year')
    shown = d2[~d2.is_hq][['year_be', 'month', 'amphoe_code']].assign(_shown=1)
    hidden = full.merge(shown, how='left', on=['year_be', 'month', 'amphoe_code'])
    hidden = hidden[hidden._shown.isna()].drop(columns='_shown')
    src = d2.drop_duplicates(['year_be', 'month'])[['year_be', 'month', 'source_file', 'source_sheet', 'loan_basis']]
    hidden = hidden.merge(src, on=['year_be', 'month']).assign(source='D2', is_hq=False, is_suppressed=True)
    d2_full = pd.concat([d2, hidden], ignore_index=True)

    d1 = read_d1(d1_rows, d1_na)
    allm = pd.concat([d2_full, d1], ignore_index=True)

    # names / province from population lookup (consistent spelling across years)
    prov_lookup = pop.drop_duplicates('amphoe_code').set_index('amphoe_code').prov_code
    allm['prov_code'] = allm.amphoe_code.map(prov_lookup)
    allm.loc[allm.is_hq, 'prov_code'] = '10'
    allm['amphoe_name'] = allm.amphoe_code.map(names).fillna(allm.get('amphoe_name'))
    allm['year_ce'] = allm.year_be - 543
    allm['period'] = allm.year_ce.astype(str) + '-' + allm.month.astype(int).astype(str).str.zfill(2)
    allm['ldr_pct'] = (allm.loan_total / allm.dep_total * 100).round(2)

    # source file/sheet per month go to a small lookup table to keep the main file compact
    allm[['period', 'source', 'loan_basis', 'source_file', 'source_sheet']].drop_duplicates().sort_values('period') \
        .to_csv(f'{out_dir}/district_monthly_sources.csv', index=False, encoding='utf-8-sig')
    for c in D1_COLS + ['dep_other']:
        allm[c] = allm[c].round().astype('Int64')          # BOT reports are in whole million baht
    allm['is_hq'] = allm.is_hq.astype(int); allm['is_suppressed'] = allm.is_suppressed.astype(int)
    cols = ['period', 'year_be', 'month', 'prov_code', 'amphoe_code', 'amphoe_name'] + D1_COLS + ['dep_other', 'ldr_pct',
            'is_hq', 'is_suppressed', 'source', 'loan_basis']
    allm = allm[cols].sort_values(['period', 'is_hq', 'amphoe_code'], ascending=[True, False, True])
    allm.to_csv(f'{out_dir}/district_monthly_D2_D1.csv', index=False, encoding='utf-8-sig')
    pd.DataFrame({'column': D1_COLS, 'label_th_in_D1': D1_TH}).to_csv(f'{out_dir}/column_dictionary.csv', index=False,
                                                                        encoding='utf-8-sig')
    pd.DataFrame(log, columns=['file', 'sheet', 'period', 'fields']).to_csv(f'{out_dir}/_parse_log.csv', index=False,
                                                                           encoding='utf-8-sig')
    return allm


if __name__ == '__main__':
    out = main(*sys.argv[1:6])
    print(out.shape)
